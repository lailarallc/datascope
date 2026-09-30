"""Tests for datascope.cli -- argument parsing, error paths, and stdout output."""

from __future__ import annotations

import textwrap
from pathlib import Path

import pytest

from datascope import __version__
from datascope.cli import main

# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

SAMPLES_DIR = Path(__file__).resolve().parent.parent / "samples" / "input"
SAMPLE_XLSX = SAMPLES_DIR / "sample_mixed_types.xlsx"
SAMPLE_SALES = SAMPLES_DIR / "sample_sales.xlsx"


def _write_csv(tmp_path: Path, text: str, name: str = "test.csv") -> Path:
    p = tmp_path / name
    p.write_text(textwrap.dedent(text).lstrip(), encoding="utf-8")
    return p


# ===================================================================
# --version
# ===================================================================

class TestVersion:

    def test_version_prints_and_exits(self, capsys):
        with pytest.raises(SystemExit) as exc_info:
            main(["--version"])
        assert exc_info.value.code == 0
        captured = capsys.readouterr()
        assert f"datascope {__version__}" in captured.out


# ===================================================================
# --help
# ===================================================================

class TestHelp:

    def test_help_prints_usage_and_exits(self, capsys):
        with pytest.raises(SystemExit) as exc_info:
            main(["--help"])
        assert exc_info.value.code == 0
        captured = capsys.readouterr()
        assert "usage:" in captured.out.lower() or "input_file" in captured.out


# ===================================================================
# Error paths
# ===================================================================

class TestErrorPaths:

    def test_no_args_exits_with_code_1(self, capsys):
        with pytest.raises(SystemExit) as exc_info:
            main([])
        assert exc_info.value.code == 1

    def test_file_not_found_exits_1(self, capsys, tmp_path):
        missing = tmp_path / "does_not_exist.xlsx"
        with pytest.raises(SystemExit) as exc_info:
            main([str(missing)])
        assert exc_info.value.code == 1
        captured = capsys.readouterr()
        assert "not found" in captured.err.lower()

    def test_unsupported_extension_exits_1(self, capsys, tmp_path):
        json_file = tmp_path / "data.json"
        json_file.write_text("{}", encoding="utf-8")
        with pytest.raises(SystemExit) as exc_info:
            main([str(json_file)])
        assert exc_info.value.code == 1
        captured = capsys.readouterr()
        assert "unsupported" in captured.err.lower()

    def test_unsupported_extension_message_lists_supported(self, capsys, tmp_path):
        json_file = tmp_path / "data.json"
        json_file.write_text("{}", encoding="utf-8")
        with pytest.raises(SystemExit):
            main([str(json_file)])
        captured = capsys.readouterr()
        assert ".csv" in captured.err
        assert ".xlsx" in captured.err

    @pytest.mark.parametrize(
        "content",
        [b"not a zip at all", b"", b"PK\x03\x04truncated"],
        ids=["garbage", "empty", "truncated"],
    )
    def test_corrupt_xlsx_exits_1_with_message_not_traceback(self, capsys, tmp_path, content):
        bad = tmp_path / "broken.xlsx"
        bad.write_bytes(content)
        with pytest.raises(SystemExit) as exc_info:
            main([str(bad)])
        assert exc_info.value.code == 1
        err = capsys.readouterr().err
        assert "broken.xlsx is not a readable .xlsx file" in err
        assert "Traceback" not in err

    def test_zip_without_workbook_exits_1_with_message(self, capsys, tmp_path):
        import zipfile

        bad = tmp_path / "notexcel.xlsx"
        with zipfile.ZipFile(bad, "w") as z:
            z.writestr("hello.txt", "x")
        with pytest.raises(SystemExit) as exc_info:
            main([str(bad)])
        assert exc_info.value.code == 1
        assert "not a readable .xlsx file" in capsys.readouterr().err


# ===================================================================
# Happy path -- xlsx
# ===================================================================

class TestHappyPathXlsx:

    def test_analyze_xlsx_creates_pdf(self, capsys, tmp_path):
        out_dir = tmp_path / "reports"
        main([str(SAMPLE_XLSX), "--output-dir", str(out_dir)])
        capsys.readouterr()

        pdf = out_dir / "sample_mixed_types_diagnostic.pdf"
        assert pdf.exists()
        assert pdf.stat().st_size > 0

    def test_stdout_contains_filename(self, capsys, tmp_path):
        out_dir = tmp_path / "reports"
        main([str(SAMPLE_XLSX), "--output-dir", str(out_dir)])
        captured = capsys.readouterr()
        assert "sample_mixed_types.xlsx" in captured.out

    def test_stdout_contains_finding_count(self, capsys, tmp_path):
        out_dir = tmp_path / "reports"
        main([str(SAMPLE_XLSX), "--output-dir", str(out_dir)])
        captured = capsys.readouterr()
        assert "Found" in captured.out
        assert "finding" in captured.out

    def test_stdout_contains_row_column_info(self, capsys, tmp_path):
        out_dir = tmp_path / "reports"
        main([str(SAMPLE_XLSX), "--output-dir", str(out_dir)])
        captured = capsys.readouterr()
        assert "200 rows" in captured.out
        assert "4 columns" in captured.out

    def test_stdout_contains_report_path(self, capsys, tmp_path):
        out_dir = tmp_path / "reports"
        main([str(SAMPLE_XLSX), "--output-dir", str(out_dir)])
        captured = capsys.readouterr()
        assert "Report saved:" in captured.out
        assert "diagnostic.pdf" in captured.out


# ===================================================================
# Happy path -- csv
# ===================================================================

class TestHappyPathCsv:

    def test_analyze_csv_creates_pdf(self, capsys, tmp_path):
        csv_path = _write_csv(tmp_path, """\
            id,name,score
            1,Alice,95
            2,Bob,N/A
            3,Carol,88
        """)
        out_dir = tmp_path / "out"
        main([str(csv_path), "--output-dir", str(out_dir)])

        pdf = out_dir / "test_diagnostic.pdf"
        assert pdf.exists()
        assert pdf.stat().st_size > 0

    def test_csv_stdout_shows_analyzing(self, capsys, tmp_path):
        csv_path = _write_csv(tmp_path, """\
            x,y
            1,2
        """)
        out_dir = tmp_path / "out"
        main([str(csv_path), "--output-dir", str(out_dir)])
        captured = capsys.readouterr()
        assert "Analyzing" in captured.out


# ===================================================================
# --output-dir creates directory
# ===================================================================

class TestOutputDir:

    def test_creates_nested_output_directory(self, capsys, tmp_path):
        csv_path = _write_csv(tmp_path, """\
            a,b
            1,2
        """)
        out_dir = tmp_path / "deep" / "nested" / "reports"
        assert not out_dir.exists()

        main([str(csv_path), "--output-dir", str(out_dir)])
        assert out_dir.exists()
        assert (out_dir / "test_diagnostic.pdf").exists()


# ===================================================================
# --sheet argument
# ===================================================================

class TestSheetArgument:

    def test_sheet_by_name(self, capsys, tmp_path):
        out_dir = tmp_path / "reports"
        main([str(SAMPLE_XLSX), "--sheet", "Sales", "--output-dir", str(out_dir)])
        captured = capsys.readouterr()
        assert "sample_mixed_types.xlsx" in captured.out
        assert (out_dir / "sample_mixed_types_diagnostic.pdf").exists()

    def test_sheet_by_index(self, capsys, tmp_path):
        out_dir = tmp_path / "reports"
        main([str(SAMPLE_XLSX), "--sheet", "0", "--output-dir", str(out_dir)])
        capsys.readouterr()
        assert (out_dir / "sample_mixed_types_diagnostic.pdf").exists()


# ===================================================================
# Empty file handling
# ===================================================================

class TestEmptyFile:

    def test_empty_csv_no_crash(self, capsys, tmp_path):
        csv_path = tmp_path / "empty.csv"
        csv_path.write_text("a,b\n", encoding="utf-8")
        out_dir = tmp_path / "reports"
        main([str(csv_path), "--output-dir", str(out_dir)])
        captured = capsys.readouterr()
        assert "Report saved:" in captured.out


# ===================================================================
# Size guard (--max-rows and cell limits)
# ===================================================================

class TestSizeGuard:

    def test_max_rows_aborts(self, tmp_path):
        csv_path = _write_csv(tmp_path, """\
            a,b
            1,2
            3,4
            5,6
        """)
        with pytest.raises(SystemExit, match="1"):
            main([str(csv_path), "--max-rows", "2", "--output-dir", str(tmp_path / "out")])

    def test_max_rows_allows_within_limit(self, capsys, tmp_path):
        csv_path = _write_csv(tmp_path, """\
            a,b
            1,2
            3,4
        """)
        out_dir = tmp_path / "out"
        main([str(csv_path), "--max-rows", "10", "--output-dir", str(out_dir)])
        assert (out_dir / "test_diagnostic.pdf").exists()

    def test_max_rows_stderr_message(self, capsys, tmp_path):
        csv_path = _write_csv(tmp_path, """\
            a,b
            1,2
            3,4
            5,6
        """)
        with pytest.raises(SystemExit):
            main([str(csv_path), "--max-rows", "1", "--output-dir", str(tmp_path / "out")])
        captured = capsys.readouterr()
        assert "--max-rows" in captured.err


# ===================================================================
# Output file naming
# ===================================================================

class TestOutputNaming:

    def test_output_filename_uses_stem(self, capsys, tmp_path):
        csv_path = _write_csv(tmp_path, """\
            x
            1
        """, name="my_data_file.csv")
        out_dir = tmp_path / "out"
        main([str(csv_path), "--output-dir", str(out_dir)])
        assert (out_dir / "my_data_file_diagnostic.pdf").exists()


# ===================================================================
# A check that crashes must not produce a report that reads as complete
# ===================================================================

class TestFailedCheckIsDisclosed:

    @pytest.fixture
    def broken_mixed_dates(self, monkeypatch):
        import datascope.analyzers as analyzers

        def boom(_result):
            raise RuntimeError("simulated analyzer crash")

        boom.__name__ = "analyze_mixed_dates"
        monkeypatch.setattr(analyzers, "analyze_mixed_dates", boom)

    _MSG = "Incomplete report: 1 check failed to run (mixed dates)"

    def test_stdout_and_json_disclose_failed_check(self, broken_mixed_dates, capsys, tmp_path):
        import json

        main([str(SAMPLE_XLSX), "--output-dir", str(tmp_path), "--format", "both"])
        out = capsys.readouterr()
        assert self._MSG in out.out
        assert "analyze_mixed_dates failed" in out.err
        payload = json.loads((tmp_path / "sample_mixed_types_diagnostic.json").read_text(encoding="utf-8"))
        assert payload["source"]["failed_checks"] == ["mixed dates"]

    def test_html_shows_banner(self, broken_mixed_dates, capsys, tmp_path):
        main([str(SAMPLE_XLSX), "--output-dir", str(tmp_path), "--format", "html"])
        html = (tmp_path / "sample_mixed_types_diagnostic.html").read_text(encoding="utf-8")
        assert 'class="incomplete" role="alert"' in html
        assert self._MSG in html

    def test_annotated_excel_shows_warning(self, broken_mixed_dates, capsys, tmp_path):
        from openpyxl import load_workbook

        main([str(SAMPLE_XLSX), "--output-dir", str(tmp_path), "--format", "annotated-excel"])
        ws = load_workbook(tmp_path / "sample_mixed_types_annotated.xlsx")["Findings"]
        texts = [c.value for c in ws["A"] if isinstance(c.value, str)]
        assert any(t.startswith(self._MSG) for t in texts)

    def test_pdf_title_page_shows_warning(self, broken_mixed_dates, capsys, tmp_path):
        pypdf = pytest.importorskip("pypdf")
        main([str(SAMPLE_XLSX), "--output-dir", str(tmp_path)])
        reader = pypdf.PdfReader(tmp_path / "sample_mixed_types_diagnostic.pdf")
        first_page = " ".join(reader.pages[0].extract_text().split())
        assert "Incomplete report" in first_page and "mixed dates" in first_page

    def test_no_failure_means_no_warning(self, capsys, tmp_path):
        main([str(SAMPLE_XLSX), "--output-dir", str(tmp_path), "--format", "html"])
        assert "Incomplete report" not in capsys.readouterr().out
        html = (tmp_path / "sample_mixed_types_diagnostic.html").read_text(encoding="utf-8")
        assert 'class="incomplete"' not in html
